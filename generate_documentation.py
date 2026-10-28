# -*- coding: utf-8 -*-
"""
Gerador da Documentação Acadêmica ABNT do FinOps Simulator.
Produz:
- Documentacao_Academica_FinOps_Simulator_ABNT.docx
- Documentacao_Academica_FinOps_Simulator_ABNT.pdf
"""

import os
import sys
import html
import subprocess
import pathlib
import io

import docx
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

import pypdf
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas

BASE_DIR = pathlib.Path(__file__).parent.resolve()

def read_code_file(rel_path):
    full_path = BASE_DIR / rel_path
    if not full_path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {full_path}")
    with open(full_path, "r", encoding="utf-8") as f:
        return f.read().strip()

print("1/5: Lendo todos os arquivos do FinOps Simulator...")

# Model
code_DatabaseHelper = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/model/DatabaseHelper.java")
code_Projeto = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/model/Projeto.java")
code_ProjetoDAO = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/model/ProjetoDAO.java")
code_TipoRecurso = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/model/TipoRecurso.java")
code_TipoRecursoDAO = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/model/TipoRecursoDAO.java")
code_RecursoAlocado = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/model/RecursoAlocado.java")
code_RecursoAlocadoDAO = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/model/RecursoAlocadoDAO.java")

# Telas e Componentes
code_activity_main = read_code_file("app/src/main/res/layout/activity_main.xml")
code_header_home = read_code_file("app/src/main/res/layout/header_home.xml")
code_item_projeto_home = read_code_file("app/src/main/res/layout/item_projeto_home.xml")
code_MainActivity = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/activity/MainActivity.java")
code_HomeProjetoAdapter = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/adapter/HomeProjetoAdapter.java")

code_activity_lista_projetos = read_code_file("app/src/main/res/layout/activity_lista_projetos.xml")
code_item_projeto = read_code_file("app/src/main/res/layout/item_projeto.xml")
code_ListaProjetosActivity = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/activity/ListaProjetosActivity.java")
code_ProjetoAdapter = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/adapter/ProjetoAdapter.java")

code_activity_form_projeto = read_code_file("app/src/main/res/layout/activity_form_projeto.xml")
code_FormProjetoActivity = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/activity/FormProjetoActivity.java")

code_activity_dashboard = read_code_file("app/src/main/res/layout/activity_dashboard.xml")
code_item_recurso = read_code_file("app/src/main/res/layout/item_recurso.xml")
code_DashboardActivity = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/activity/DashboardActivity.java")
code_RecursoAdapter = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/adapter/RecursoAdapter.java")

code_activity_form_recurso = read_code_file("app/src/main/res/layout/activity_form_recurso.xml")
code_FormRecursoActivity = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/activity/FormRecursoActivity.java")

code_activity_lista_tipos = read_code_file("app/src/main/res/layout/activity_lista_tipos.xml")
code_item_tipo_recurso = read_code_file("app/src/main/res/layout/item_tipo_recurso.xml")
code_ListaTiposActivity = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/activity/ListaTiposActivity.java")
code_TipoRecursoAdapter = read_code_file("app/src/main/java/com/lucasjsm/finops_simulator/adapter/TipoRecursoAdapter.java")

# Utilitários e Recursos
code_bottom_nav_menu = read_code_file("app/src/main/res/menu/bottom_nav_menu.xml")
code_strings = read_code_file("app/src/main/res/values/strings.xml")
code_colors = read_code_file("app/src/main/res/values/colors.xml")
code_themes = read_code_file("app/src/main/res/values/themes.xml")

code_bg_header_blue = read_code_file("app/src/main/res/drawable/bg_header_blue.xml")
code_bg_hero_card = read_code_file("app/src/main/res/drawable/bg_hero_card.xml")
code_bg_button_primary = read_code_file("app/src/main/res/drawable/bg_button_primary.xml")
code_bg_button_secondary = read_code_file("app/src/main/res/drawable/bg_button_secondary.xml")
code_bg_button_white = read_code_file("app/src/main/res/drawable/bg_button_white.xml")
code_bg_chip_cost = read_code_file("app/src/main/res/drawable/bg_chip_cost.xml")
code_bg_chip_cost_warning = read_code_file("app/src/main/res/drawable/bg_chip_cost_warning.xml")
code_bg_chip_cost_danger = read_code_file("app/src/main/res/drawable/bg_chip_cost_danger.xml")
code_bg_edittext = read_code_file("app/src/main/res/drawable/bg_edittext.xml")
code_bg_spinner_wrapper = read_code_file("app/src/main/res/drawable/bg_spinner_wrapper.xml")
code_bg_custo_estimado = read_code_file("app/src/main/res/drawable/bg_custo_estimado.xml")
code_progress_budget = read_code_file("app/src/main/res/drawable/progress_budget.xml")
code_progress_budget_warning = read_code_file("app/src/main/res/drawable/progress_budget_warning.xml")
code_progress_budget_danger = read_code_file("app/src/main/res/drawable/progress_budget_danger.xml")

code_manifest = read_code_file("app/src/main/AndroidManifest.xml")

print("2/5: Construindo documento DOCX...")

doc = Document()

# Configuração de Margens ABNT NBR 14724
sec1 = doc.sections[0]
sec1.top_margin = Cm(3.0)
sec1.left_margin = Cm(3.0)
sec1.bottom_margin = Cm(2.0)
sec1.right_margin = Cm(2.0)
sec1.page_width = Cm(21.0)
sec1.page_height = Cm(29.7)

style_normal = doc.styles['Normal']
font_normal = style_normal.font
font_normal.name = 'Arial'
font_normal.size = Pt(12)
font_normal.color.rgb = RGBColor(0, 0, 0)

def add_title_p(doc, text, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    return p

def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    return p

def add_bullet_p(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run("- " + text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(13)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    return p

def add_placeholder_screenshot(doc, label):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    run = p.add_run(label)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    return p

def add_code_block(doc, filepath, code):
    p_file = doc.add_paragraph()
    p_file.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_file.paragraph_format.space_before = Pt(12)
    p_file.paragraph_format.space_after = Pt(2)
    p_file.paragraph_format.keep_with_next = True
    run_file = p_file.add_run(f"Arquivo: {filepath}")
    run_file.bold = True
    run_file.font.name = 'Arial'
    run_file.font.size = Pt(10)

    lines = code.splitlines()
    for line in lines:
        p_code = doc.add_paragraph()
        p_code.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_code.paragraph_format.space_before = Pt(0)
        p_code.paragraph_format.space_after = Pt(0)
        p_code.paragraph_format.line_spacing = 1.0
        run_code = p_code.add_run(line if line else " ")
        run_code.font.name = 'Consolas'
        run_code.font.size = Pt(8.5)
        run_code.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)

# ==================== CAPA ====================
add_title_p(doc, "UNIVERSIDADE DO ESTADO DO RIO DE JANEIRO — UERJ", bold=True, size=12, space_before=10)
add_title_p(doc, "FACULDADE DE ENGENHARIA — FEN", bold=True, size=12)
add_title_p(doc, "DEPARTAMENTO DE ENGENHARIA DE SISTEMAS E COMPUTAÇÃO", bold=True, size=12, space_after=90)

add_title_p(doc, "[INSERIR NOME COMPLETO DO ALUNO — MATRÍCULA: XXXXXXXXX]", bold=True, size=12, space_after=120)

add_title_p(doc, "FINOPS SIMULATOR: SISTEMA MÓVEL DE SIMULAÇÃO E GESTÃO DE CUSTOS DE NUVEM COMPUTACIONAL", bold=True, size=14, space_after=160)

add_title_p(doc, "Rio de Janeiro", bold=False, size=12)
add_title_p(doc, "2026", bold=False, size=12)

doc.add_page_break()

# ==================== FOLHA DE ROSTO ====================
add_title_p(doc, "FOLHA DE ROSTO (CONTRACAPA)", bold=True, size=12, space_before=10, space_after=40)
add_title_p(doc, "[INSERIR NOME COMPLETO DO ALUNO — MATRÍCULA: XXXXXXXXX]", bold=True, size=12, space_after=80)

add_title_p(doc, "FINOPS SIMULATOR: SISTEMA MÓVEL DE SIMULAÇÃO E GESTÃO DE CUSTOS DE NUVEM COMPUTACIONAL", bold=True, size=14, space_after=60)

p_nat = doc.add_paragraph()
p_nat.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_nat.paragraph_format.left_indent = Cm(8.0)
p_nat.paragraph_format.space_before = Pt(20)
p_nat.paragraph_format.space_after = Pt(14)
p_nat.paragraph_format.line_spacing = 1.0
run_nat = p_nat.add_run("Trabalho acadêmico apresentado à disciplina de Programação para Dispositivos Móveis, do Curso de Engenharia de Sistemas e Computação da Universidade do Estado do Rio de Janeiro (UERJ), como requisito parcial de avaliação.")
run_nat.font.name = 'Arial'
run_nat.font.size = Pt(10)

p_prof = doc.add_paragraph()
p_prof.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_prof.paragraph_format.left_indent = Cm(8.0)
p_prof.paragraph_format.space_before = Pt(10)
p_prof.paragraph_format.space_after = Pt(100)
p_prof.paragraph_format.line_spacing = 1.0
run_prof = p_prof.add_run("Professor Orientador: Prof. Dr. Engº Carlos A. Sicsú A. do Nascimento")
run_prof.bold = True
run_prof.font.name = 'Arial'
run_prof.font.size = Pt(10)

add_title_p(doc, "Rio de Janeiro", bold=False, size=12)
add_title_p(doc, "2026", bold=False, size=12)

doc.add_page_break()

# ==================== SUMÁRIO ====================
add_title_p(doc, "SUMÁRIO (ÍNDICE)", bold=True, size=14, space_before=10, space_after=24)

sumario_linhas = [
    ("1. INTRODUÇÃO", True, 0),
    ("- 1.1 Contextualização e Tema do Projeto", False, 0.5),
    ("- 1.2 Objetivos Geral e Específicos", False, 0.5),
    ("- 1.3 Padrão Arquitetural MVC e Persistência SQLite", False, 0.5),
    ("2. ARQUITETURA DO SISTEMA E MODELAGEM DE DADOS", True, 0),
    ("- 2.1 Divisão em Pacotes segundo o Padrão MVC (model, view, controller)", False, 0.5),
    ("- 2.2 Esquema Relacional do Banco de Dados SQLite (finops_simulator.db)", False, 0.5),
    ("- 2.3 Regras de Negócio e Integridade de Domínio", False, 0.5),
    ("3. CAMADA MODEL — CLASSES DE ENTIDADE E PERSISTÊNCIA (DAO)", True, 0),
    ("- 3.1 Gerenciador do Banco de Dados (DatabaseHelper.java)", False, 0.5),
    ("- 3.2 Entidade e Persistência de Projeto (Projeto.java e ProjetoDAO.java)", False, 0.5),
    ("- 3.3 Entidade e Persistência de Tipo de Recurso (TipoRecurso.java e TipoRecursoDAO.java)", False, 0.5),
    ("- 3.4 Entidade e Persistência de Recurso Alocado (RecursoAlocado.java e RecursoAlocadoDAO.java)", False, 0.5),
    ("4. APRESENTAÇÃO DAS TELAS E CÓDIGOS-FONTE (VIEW E CONTROLLER)", True, 0),
    ("- 4.1 Tela Principal — Visão Geral e Acesso Rápido (MainActivity, HomeProjetoAdapter, activity_main.xml, header_home.xml, item_projeto_home.xml)", False, 0.5),
    ("- 4.2 Tela de Lista de Projetos (ListaProjetosActivity, ProjetoAdapter, activity_lista_projetos.xml, item_projeto.xml)", False, 0.5),
    ("- 4.3 Tela de Cadastro e Edição de Projeto (FormProjetoActivity, activity_form_projeto.xml)", False, 0.5),
    ("- 4.4 Tela de Dashboard e Gestão do Projeto (DashboardActivity, RecursoAdapter, activity_dashboard.xml, item_recurso.xml)", False, 0.5),
    ("- 4.5 Tela de Alocação e Edição de Recurso (FormRecursoActivity, activity_form_recurso.xml)", False, 0.5),
    ("- 4.6 Tela de Catálogo de Serviços de Nuvem (ListaTiposActivity, TipoRecursoAdapter, activity_lista_tipos.xml, item_tipo_recurso.xml)", False, 0.5),
    ("- 4.7 Utilitários de Interface, Recursos XML e Manifesto (bottom_nav_menu.xml, strings.xml, colors.xml, themes.xml, Drawables e AndroidManifest.xml)", False, 0.5),
    ("5. CONCLUSÃO", True, 0),
    ("6. REFERÊNCIAS BIBLIOGRÁFICAS", True, 0),
]

for texto, bold, indent in sumario_linhas:
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.left_indent = Cm(indent)
    p_sum.paragraph_format.space_before = Pt(3 if bold else 1)
    p_sum.paragraph_format.space_after = Pt(2)
    p_sum.paragraph_format.line_spacing = 1.3
    run_sum = p_sum.add_run(texto)
    run_sum.bold = bold
    run_sum.font.name = 'Arial'
    run_sum.font.size = Pt(11)

# Seção 2: Início do Texto e Numeração de Página
sec2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
sec2.top_margin = Cm(3.0)
sec2.left_margin = Cm(3.0)
sec2.bottom_margin = Cm(2.0)
sec2.right_margin = Cm(2.0)
sec2.header.is_linked_to_previous = False

header = sec2.header
p_head = header.paragraphs[0]
p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run_page_num = p_head.add_run()
fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
run_page_num._r.append(fldChar1)
run_page_num._r.append(instrText)
run_page_num._r.append(fldChar2)
run_page_num._r.append(fldChar3)
run_page_num.font.name = 'Arial'
run_page_num.font.size = Pt(10)

# ==================== 1. INTRODUÇÃO ====================
add_heading_1(doc, "1. INTRODUÇÃO")

add_heading_2(doc, "1.1 Contextualização e Tema do Projeto")
add_body_p(doc, "A crescente adoção do paradigma de computação em nuvem (Cloud Computing) transformou profundamente a arquitetura, o provisionamento e a operação de sistemas de software corporativos modernos. A migração de ambientes locais físicos (on-premises) para os principais provedores de nuvem pública do mercado global — notadamente Amazon Web Services (AWS), Google Cloud Platform (GCP), Microsoft Azure e Oracle Cloud Infrastructure (OCI) — proporcionou às organizações elasticidade operacional incomparável, alta disponibilidade regional e capacidade de rápido escalonamento de aplicações. Todavia, a transição para um modelo financeiro puramente variável baseado no consumo efetivo sob demanda (pay-as-you-go) trouxe consigo novos e severos desafios de previsibilidade orçamentária, tornando comum o surgimento de gastos imprevistos e estouros recorrentes de orçamento (budget overruns).")
add_body_p(doc, "Nesse contexto de complexidade operacional e financeira, consolidou-se internacionalmente a disciplina de Cloud FinOps (Financial Operations em Nuvem), padronizada e fomentada pela FinOps Foundation. O FinOps representa uma evolução cultural e técnica interdisciplinar que combina práticas financeiras tradicionais, engenharia de confiabilidade e operações DevOps, capacitando equipes de desenvolvimento e engenharia a assumirem a responsabilidade financeira direta pelos recursos computacionais que provisionam. Seu objetivo fundamental não é apenas conter custos de maneira indiscriminada, mas sim viabilizar decisões arquiteturais fundamentadas em valor de negócio, garantindo máxima eficiência por unidade monetária investida em infraestrutura.")
add_body_p(doc, "Com o propósito de aplicar e consolidar esses conceitos em um ambiente prático na plataforma móvel Android, foi concebido e implementado o aplicativo FinOps Simulator (desenvolvido integralmente em linguagem Java). O sistema atua como um simulador estruturado de custos e ferramenta móvel de governança orçamentária multicloud, voltado a engenheiros de software, arquitetos de nuvem, gestores de tecnologia e estudantes universitários.")
add_body_p(doc, "O aplicativo permite cadastrar e manter projetos independentes associados às quatro maiores plataformas de nuvem do mercado (AWS, GCP, Azure e OCI), com estipulação de teto orçamentário mensal (budget). A partir de um catálogo pré-carregado contendo 14 serviços emblemáticos de infraestrutura (como instâncias elásticas de computação, armazenamento de objetos, data warehouses analíticos de alto desempenho, bancos de dados NoSQL e clusters Kubernetes gerenciados), o usuário pode dimensionar e alocar recursos a cada projeto. O simulador apura de maneira instantânea o custo financeiro estimado com base nas métricas reais de cobrança (incluindo o fator temporal de 730 horas mensais para serviços contínuos), apresentando em tempo real o percentual de consumo do orçamento e sinalizando preventivamente alertas visuais cromáticos (faixas de conformidade, atenção a partir de 80% e perigo/estouro acima de 100%).")

add_heading_2(doc, "1.2 Objetivos Geral e Específicos")
add_body_p(doc, "O objetivo geral deste trabalho acadêmico consiste em projetar, arquitetar e implementar uma aplicação Android nativa e completa, desenvolvida em linguagem Java, adotando o padrão arquitetural MVC (Model-View-Controller) e utilizando persistência local em banco de dados relacional SQLite, consolidando na prática as competências e metodologias fundamentais de desenvolvimento para dispositivos móveis ministradas na disciplina.")
add_body_p(doc, "Como objetivos específicos que nortearam o desenvolvimento do projeto, destacam-se:")
add_bullet_p(doc, "Implementar operações completas de CRUD (Create, Read, Update, Delete) para as três entidades nucleares do sistema: Projeto, Tipo de Recurso (Catálogo) e Recurso Alocado (tabela associativa);")
add_bullet_p(doc, "Isolar com rigor as regras de negócio, a lógica matemática de precificação em nuvem (fator de 730 horas/mês, conversões de unidade e somatórios agregados) e as manipulações relacionais SQL na camada Model (com.lucasjsm.finops_simulator.model);")
add_bullet_p(doc, "Construir interfaces gráficas responsivas, intuitivas e aderentes ao Material Design na camada View (res/layout e com.lucasjsm.finops_simulator.adapter), utilizando ListView associado ao padrão de reciclagem ViewHolder, MaterialCardView, ProgressBar dinâmico, ScrollView, Spinner e BottomNavigationView;")
add_bullet_p(doc, "Implementar filtros combinados e reativos no Catálogo de Serviços, cálculo instantâneo em tempo real de custo (live preview) nos formulários de cadastro através do padrão TextWatcher, e agregação dinâmica automática de custos no painel de controle;")
add_bullet_p(doc, "Assegurar a integridade relacional e a consistência histórica do banco de dados SQLite, implementando exclusão em cascata (cascade delete) para eliminar automaticamente os recursos vinculados ao remover um projeto, além de salvaguardar regras de domínio como a imutabilidade do provedor do projeto e o bloqueio de serviço na alteração de alocações.")

add_heading_2(doc, "1.3 Padrão Arquitetural MVC e Persistência SQLite")
add_body_p(doc, "Em rigorosa conformidade com o programa curricular e a metodologia pedagógica apresentada em aula pelo Prof. Dr. Engº Carlos A. Sicsú A. do Nascimento, o FinOps Simulator adota a separação clássica de responsabilidades em três camadas fundamentais:")
add_bullet_p(doc, "Model (com.lucasjsm.finops_simulator.model): Contém as classes de entidade de domínio (Projeto, TipoRecurso e RecursoAlocado), a classe centralizadora de gerenciamento do banco de dados SQLite (DatabaseHelper, subclasse de SQLiteOpenHelper) e as classes especializadas de acesso e persistência a dados (ProjetoDAO, TipoRecursoDAO e RecursoAlocadoDAO). Toda a lógica financeira de cálculo de custo de nuvem, conversão de horas operacionais, somatórios agregados, consultas parametrizadas com filtros e operações transacionais com exclusão em cascata residem exclusivamente nesta camada, sem qualquer dependência das bibliotecas visuais do Android.")
add_bullet_p(doc, "View (res/layout/*.xml e com.lucasjsm.finops_simulator.adapter): Compreende a totalidade dos arquivos XML declarativos de layout, drawables vetoriais customizados, paleta semântica de cores, temas visuais e as quatro classes de adaptadores de lista (HomeProjetoAdapter, ProjetoAdapter, RecursoAdapter e TipoRecursoAdapter). Os adaptadores estendem ArrayAdapter e utilizam o padrão estático ViewHolder para evitar chamadas redundantes a findViewById(), sendo responsáveis estritamente pela renderização dos dados fornecidos pelo Model, estilização dinâmica dos indicadores visuais de orçamento e disparo de eventos de interface (através da interface OnRecursoActionListener e de cliques convencionais).")
add_bullet_p(doc, "Controller (com.lucasjsm.finops_simulator.activity): Composto pelas seis Activities do aplicativo (MainActivity, ListaProjetosActivity, FormProjetoActivity, DashboardActivity, FormRecursoActivity e ListaTiposActivity). As Activities atuam exclusivamente como mediadoras e coordenadoras de fluxo: capturam os eventos e interações do usuário disparados na View, coordenam o ciclo de vida do Android (onCreate, onResume, onPause), recebem e repassam dados via Intent (utilizando parâmetros padronizados como acao = -1 para inclusão e acao = 0 para alteração), invocam as rotinas de persistência e consulta dos DAOs e atualizam a View com os dados retornados.")

# ==================== 2. ARQUITETURA ====================
doc.add_page_break()
add_heading_1(doc, "2. ARQUITETURA DO SISTEMA E MODELAGEM DE DADOS")

add_heading_2(doc, "2.1 Divisão em Pacotes segundo o Padrão MVC")
add_body_p(doc, "A estrutura de código-fonte do projeto FinOps Simulator está organizada de maneira modularizada e coesa nos seguintes pacotes Java:")
add_bullet_p(doc, "com.lucasjsm.finops_simulator.model: Abriga as classes DatabaseHelper, Projeto, ProjetoDAO, TipoRecurso, TipoRecursoDAO, RecursoAlocado e RecursoAlocadoDAO;")
add_bullet_p(doc, "com.lucasjsm.finops_simulator.adapter: Abriga as classes HomeProjetoAdapter, ProjetoAdapter, RecursoAdapter e TipoRecursoAdapter, complementadas pelas definições visuais em res/layout/ e drawables em res/drawable/;")
add_bullet_p(doc, "com.lucasjsm.finops_simulator.activity: Abriga as classes controladoras MainActivity, ListaProjetosActivity, FormProjetoActivity, DashboardActivity, FormRecursoActivity e ListaTiposActivity.")

add_heading_2(doc, "2.2 Esquema Relacional do Banco de Dados SQLite (finops_simulator.db)")
add_body_p(doc, "O banco de dados local da aplicação, denominado finops_simulator.db (versão 1), é inteiramente administrado pela classe DatabaseHelper. O esquema relacional é estruturado em três tabelas associadas por integridade referencial:")
add_bullet_p(doc, "1. Tabela projeto: Armazena os dados cadastrais e orçamentários dos projetos simulados. Possui os campos id (INTEGER PRIMARY KEY AUTOINCREMENT), nome (TEXT NOT NULL), descricao (TEXT), provedor (TEXT NOT NULL), orcamento_maximo (REAL) e data_criacao (TEXT). Apresenta relacionamento de cardinalidade 1 para N com a tabela recurso_alocado.")
add_bullet_p(doc, "2. Tabela tipo_recurso: Representa o catálogo pré-cadastrado de serviços de nuvem computacional. Possui as colunas id (INTEGER PRIMARY KEY AUTOINCREMENT), nome (TEXT NOT NULL), provedor (TEXT NOT NULL), unidade (TEXT NOT NULL), custo_unitario (REAL NOT NULL) e descricao (TEXT). No momento da criação do banco, são automaticamente injetados 14 registros oficiais de serviços das quatro maiores plataformas globais (AWS, GCP, Azure e OCI), cobrindo categorias de computação, armazenamento, bancos de dados e analytics.")
add_bullet_p(doc, "3. Tabela recurso_alocado: Tabela associativa que relaciona um projeto a um tipo de recurso dimensionado, viabilizando o relacionamento relacional N para N entre as entidades através de uma entidade associativa com atributos próprios. Contém os campos id (INTEGER PRIMARY KEY AUTOINCREMENT), projeto_id (INTEGER NOT NULL, chave estrangeira para projeto.id), tipo_recurso_id (INTEGER NOT NULL, chave estrangeira para tipo_recurso.id), quantidade (REAL NOT NULL), periodo (REAL NOT NULL), custo_estimado (REAL) e observacao (TEXT).")

add_heading_2(doc, "2.3 Regras de Negócio e Integridade de Domínio")
add_body_p(doc, "A integridade lógica do FinOps Simulator é assegurada por regras de negócio rigorosas e proteções de domínio incorporadas na camada Model:")
add_bullet_p(doc, "Cálculo de Custo Estimado com Fator de Equivalência Temporal (730 h/mês): No método calcularCusto(custoUnitario, unidadeRecurso) da classe RecursoAlocado, o sistema implementa a convenção oficial da indústria de computação em nuvem para serviços cobrados por hora contínua. Como o ano civil médio possui 365 dias, divididos por 12 meses e multiplicados por 24 horas/dia, obtém-se o fator padronizado de 730 horas operacionais por mês. Caso a unidade seja 'hora', o multiplicador de tempo adotado é 730.0; para grandezas pontuais (como GB, TB ou milhões de requisições), o multiplicador é 1.0. A fórmula implementada é: custoEstimado = quantidade * custoUnitario * multiplicadorTempo * periodo.")
add_bullet_p(doc, "Imutabilidade do Provedor de Nuvem no Projeto: Ao criar um projeto, o provedor (AWS, GCP, Azure ou OCI) é fixado. Caso o usuário pudesse alterar o provedor de um projeto em andamento, os recursos de infraestrutura já alocados (ex.: instâncias AWS EC2) passariam a constar sob outro provedor (ex.: Google Cloud), gerando uma inconsistência de domínio inaceitável. Dessa forma, na tela de edição (FormProjetoActivity), o componente Spinner de seleção de provedor é desabilitado (spProvedor.setEnabled(false)), e o método ProjetoDAO.alterar() atualiza exclusivamente os campos nome, descricao e orcamento_maximo.")
add_bullet_p(doc, "Isolamento Contextual de Serviços por Provedor: Ao alocar um recurso no projeto (FormRecursoActivity), o catálogo de serviços apresentado para seleção é filtrado de maneira estrita pelo provedor do projeto através de tipoDao.getAllPorProvedor(projeto.getProvedor()), impedindo que serviços da AWS sejam indevidamente alocados em projetos da Azure ou OCI.")
add_bullet_p(doc, "Bloqueio do Tipo de Recurso na Edição de Alocação: Quando uma alocação de recurso já existente é editada, o tipo de recurso associado permanece fixo (spTipo.setEnabled(false)), permitindo ao usuário alterar o volume, o período operacional em meses e as observações textuais, disparando o recálculo automático do custo estimado com garantia de rastreabilidade.")
add_bullet_p(doc, "Exclusão em Cascata (Cascade Delete) no SQLite: Ao excluir um projeto, o método ProjetoDAO.apagar(long id) executa a deleção prévia de todos os registros na tabela recurso_alocado cujo projeto_id coincida com o identificador do projeto antes de remover o próprio registro da tabela projeto, impedindo a existência de registros órfãos e mantendo a sanidade relacional do banco de dados local.")
add_bullet_p(doc, "Cálculo Consolidado e Prevenção Visual de Estouro Orçamentário: O método recursoDao.calcularCustoTotal(long projetoId) recupera todos os recursos do projeto e computa o somatório de custoEstimado. Com base na proporção entre o custo total e o orçamento estipulado (pct = (custo / orcamento) * 100), as telas aplicam estilos cromáticos reativos de sinalização: Normal (< 80%, cores primárias e tons de verde saudável), Alerta (80% a 100%, tons em amarelo/âmbar com texto #B45309) e Estouro/Perigo (> 100%, tons em carmim/vermelho com texto #BE123C), alertando o usuário em tempo real sobre a necessidade de readequação de gastos.")

# ==================== 3. CAMADA MODEL ====================
doc.add_page_break()
add_heading_1(doc, "3. CAMADA MODEL — CLASSES DE ENTIDADE E PERSISTÊNCIA (DAO)")
add_body_p(doc, "A seguir são apresentados todos os códigos-fonte Java pertencentes ao pacote com.lucasjsm.finops_simulator.model, responsáveis pela integridade do esquema SQLite, representação das entidades de domínio e operações de persistência e cálculo.")

add_heading_2(doc, "3.1 Gerenciador do Banco de Dados (DatabaseHelper.java)")
add_body_p(doc, "Classe responsável pela instanciação do banco de dados finops_simulator.db, criação das tabelas projeto, tipo_recurso e recurso_alocado com suas respectivas restrições de integridade referencial, e injeção do catálogo inicial com 14 serviços de nuvem computacional.")
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/model/DatabaseHelper.java", code_DatabaseHelper)

add_heading_2(doc, "3.2 Entidade e Persistência de Projeto (Projeto.java e ProjetoDAO.java)")
add_body_p(doc, "A classe de entidade Projeto encapsula os atributos de um projeto de nuvem (nome, descrição, provedor, teto orçamentário mensal e data de criação), além do custo consolidado calculado. A classe ProjetoDAO concentra as operações de CRUD, consultas parametrizadas ordenadas por nome, atualização com proteção de imutabilidade do provedor e rotina de exclusão em cascata.")
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/model/Projeto.java", code_Projeto)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/model/ProjetoDAO.java", code_ProjetoDAO)

add_heading_2(doc, "3.3 Entidade e Persistência de Tipo de Recurso (TipoRecurso.java e TipoRecursoDAO.java)")
add_body_p(doc, "A classe TipoRecurso modela as especificações de um serviço do catálogo de nuvem (nome, provedor, unidade de medida, custo unitário e descrição técnica). A classe TipoRecursoDAO realiza as consultas ao catálogo, incluindo a listagem geral ordenada e a consulta parametrizada com cláusula WHERE filtrada estritamente pelo provedor de nuvem informado.")
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/model/TipoRecurso.java", code_TipoRecurso)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/model/TipoRecursoDAO.java", code_TipoRecursoDAO)

add_heading_2(doc, "3.4 Entidade e Persistência de Recurso Alocado (RecursoAlocado.java e RecursoAlocadoDAO.java)")
add_body_p(doc, "A classe RecursoAlocado modela a alocação de um serviço específico a um projeto, encapsulando a regra de negócio do cálculo de custo estimado baseado na métrica da unidade (aplicando o fator de 730 horas/mês para serviços contínuos). A classe RecursoAlocadoDAO implementa a persistência da tabela associativa, listagem de recursos vinculados a um projeto e cálculo agregado do custo total do projeto.")
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/model/RecursoAlocado.java", code_RecursoAlocado)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/model/RecursoAlocadoDAO.java", code_RecursoAlocadoDAO)

# ==================== 4. APRESENTAÇÃO DAS TELAS ====================
doc.add_page_break()
add_heading_1(doc, "4. APRESENTAÇÃO DAS TELAS E CÓDIGOS-FONTE (VIEW E CONTROLLER)")
add_body_p(doc, "Neste capítulo são detalhadas todas as telas da aplicação FinOps Simulator. Para cada interface gráfica, apresentam-se a descrição funcional completa, o espaço reservado para a inserção da captura de tela (print) e a totalidade dos códigos-fonte associados em XML (camada View) e Java (camadas Controller/Activity e View/Adapter).")

# 4.1 MainActivity
doc.add_page_break()
add_heading_2(doc, "4.1 Tela Principal — Visão Geral e Acesso Rápido (MainActivity)")
add_body_p(doc, "Descrição Funcional: Tela inicial do aplicativo FinOps Simulator. Apresenta o cabeçalho institucional do sistema com ícone temático de nuvem, um Hero Card estilizado em gradiente azul com atalho rápido para criação de novos projetos, uma seção informativa de seleção e a listagem dinâmica de projetos cadastrados em um ListView. Cada item da lista exibe o número sequencial do projeto, nome, provedor de nuvem, orçamento mensal configurado, chip cromático com o custo total mensal e o percentual de consumo do orçamento (com formatação condicional em verde, amarelo para >= 80% e vermelho para > 100%), além de dois botões de ação rápida: 'Abrir Dashboard' e '+ Recursos'. A navegação inferior é provida pelo BottomNavigationView padronizado.")
add_placeholder_screenshot(doc, "[INSERIR AQUI O PRINT DA TELA 1: TELA PRINCIPAL — VISÃO GERAL (MainActivity)]")
add_code_block(doc, "app/src/main/res/layout/activity_main.xml", code_activity_main)
add_code_block(doc, "app/src/main/res/layout/header_home.xml", code_header_home)
add_code_block(doc, "app/src/main/res/layout/item_projeto_home.xml", code_item_projeto_home)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/activity/MainActivity.java", code_MainActivity)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/adapter/HomeProjetoAdapter.java", code_HomeProjetoAdapter)

# 4.2 ListaProjetosActivity
doc.add_page_break()
add_heading_2(doc, "4.2 Tela de Lista de Projetos (ListaProjetosActivity)")
add_body_p(doc, "Descrição Funcional: Tela dedicada à gestão e visualização abrangente de todos os projetos cadastrados no simulador. Possui cabeçalho superior estilizado em azul com botão circular de retorno, subtítulo descritivo, ListView com divisor transparente e espaçamento harmônico entre cards, estado vazio amigável quando não há projetos registrados e botão de ação flutuante (FAB) para cadastro de novo projeto. Cada card exibe o nome do projeto, tag do provedor, custo estimado total mensal, teto orçamentário estipulado e uma barra de progresso horizontal (ProgressBar) com preenchimento cromático dinâmico que reflete o percentual de consumo do orçamento.")
add_placeholder_screenshot(doc, "[INSERIR AQUI O PRINT DA TELA 2: TELA DE LISTA DE PROJETOS (ListaProjetosActivity)]")
add_code_block(doc, "app/src/main/res/layout/activity_lista_projetos.xml", code_activity_lista_projetos)
add_code_block(doc, "app/src/main/res/layout/item_projeto.xml", code_item_projeto)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/activity/ListaProjetosActivity.java", code_ListaProjetosActivity)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/adapter/ProjetoAdapter.java", code_ProjetoAdapter)

# 4.3 FormProjetoActivity
doc.add_page_break()
add_heading_2(doc, "4.3 Tela de Cadastro e Edição de Projeto (FormProjetoActivity)")
add_body_p(doc, "Descrição Funcional: Formulário versátil utilizado tanto para a criação de um novo projeto (quando recebe acao = -1 via Intent) quanto para a edição de um projeto existente (quando recebe acao = 0 e o identificador do projeto). Contém campos para Nome do Projeto, Descrição, seletor de Provedor de Nuvem via Spinner (AWS, GCP, Azure, OCI) e campo numérico decimal para estipulação do Orçamento Máximo mensal em reais. Na modalidade de edição, o seletor de provedor é automaticamente bloqueado (spProvedor.setEnabled(false)) para resguardar a integridade dos serviços já alocados, e o botão de persistência aciona a rotina de atualização correspondente no ProjetoDAO.")
add_placeholder_screenshot(doc, "[INSERIR AQUI O PRINT DA TELA 3: TELA DE CADASTRO E EDIÇÃO DE PROJETO (FormProjetoActivity)]")
add_code_block(doc, "app/src/main/res/layout/activity_form_projeto.xml", code_activity_form_projeto)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/activity/FormProjetoActivity.java", code_FormProjetoActivity)

# 4.4 DashboardActivity
doc.add_page_break()
add_heading_2(doc, "4.4 Tela de Dashboard e Gestão do Projeto (DashboardActivity)")
add_body_p(doc, "Descrição Funcional: Painel analítico de controle de um projeto específico. Apresenta cabeçalho estilizado, card de síntese financeira com nome, provedor, teto orçamentário, custo total consumido e barra de progresso reativa com rótulo percentual. Na seção inferior, lista os recursos alocados àquele projeto através de um ListView customizado, com indicação de serviço, volume dimensionado, período em meses, custo parcial estimado e botão de exclusão individual. Na extremidade inferior direita, disponibiliza botões de ação empilhados para exclusão integral do projeto (com diálogo de confirmação AlertDialog e exclusão em cascata de recursos), edição dos dados do projeto e adição de novos recursos computacionais.")
add_placeholder_screenshot(doc, "[INSERIR AQUI O PRINT DA TELA 4: TELA DE DASHBOARD E GESTÃO DO PROJETO (DashboardActivity)]")
add_code_block(doc, "app/src/main/res/layout/activity_dashboard.xml", code_activity_dashboard)
add_code_block(doc, "app/src/main/res/layout/item_recurso.xml", code_item_recurso)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/activity/DashboardActivity.java", code_DashboardActivity)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/adapter/RecursoAdapter.java", code_RecursoAdapter)

# 4.5 FormRecursoActivity
doc.add_page_break()
add_heading_2(doc, "4.5 Tela de Alocação e Edição de Recurso (FormRecursoActivity)")
add_body_p(doc, "Descrição Funcional: Formulário responsável pelo dimensionamento de recursos de infraestrutura vinculados a um projeto. Suporta os modos de inclusão (acao = -1) e alteração (acao = 0). O seletor Spinner de tipo de recurso carrega exclusivamente os serviços pertinentes ao provedor do projeto em questão. Apresenta campos de entrada para Quantidade/Volume (em GB, instâncias ou TB), Período operacional em meses e campo opcional para Observações. Através de ouvintes TextWatcher associados aos campos de digitação e do ouvinte de seleção do Spinner, a tela atualiza em tempo real uma caixa de pré-visualização (live preview) do custo estimado resultante da alocação, aplicando o fator de 730 horas/mês quando o serviço for de tarifação horária.")
add_placeholder_screenshot(doc, "[INSERIR AQUI O PRINT DA TELA 5: TELA DE ALOCAÇÃO E EDIÇÃO DE RECURSO (FormRecursoActivity)]")
add_code_block(doc, "app/src/main/res/layout/activity_form_recurso.xml", code_activity_form_recurso)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/activity/FormRecursoActivity.java", code_FormRecursoActivity)

# 4.6 ListaTiposActivity
doc.add_page_break()
add_heading_2(doc, "4.6 Tela de Catálogo de Serviços de Nuvem (ListaTiposActivity)")
add_body_p(doc, "Descrição Funcional: Catálogo de referência dos serviços de computação em nuvem pré-configurados no banco de dados. Apresenta cabeçalho informativo, seletor Spinner de filtragem com as opções 'Todos os Provedores', 'AWS', 'GCP', 'Azure' e 'OCI', e uma listagem em ListView contendo os cards individuais de serviços. Cada card apresenta o nome do serviço, chip com coloração temática correspondente ao provedor (âmbar para AWS, azul royal para GCP, azul celeste para Azure e carmim para OCI), custo unitário formatado com precisão de três casas decimais por unidade métrica e descrição técnica de finalidade. Ao tocar em um item da lista, exibe uma mensagem Toast informativa com a especificação consolidada do serviço.")
add_placeholder_screenshot(doc, "[INSERIR AQUI O PRINT DA TELA 6: TELA DE CATÁLOGO DE SERVIÇOS DE NUVEM (ListaTiposActivity)]")
add_code_block(doc, "app/src/main/res/layout/activity_lista_tipos.xml", code_activity_lista_tipos)
add_code_block(doc, "app/src/main/res/layout/item_tipo_recurso.xml", code_item_tipo_recurso)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/activity/ListaTiposActivity.java", code_ListaTiposActivity)
add_code_block(doc, "app/src/main/java/com/lucasjsm/finops_simulator/adapter/TipoRecursoAdapter.java", code_TipoRecursoAdapter)

# 4.7 Utilitários de Interface, Recursos XML e Manifesto
doc.add_page_break()
add_heading_2(doc, "4.7 Utilitários de Interface, Recursos XML e Manifesto")
add_body_p(doc, "Abaixo encontram-se o menu de navegação inferior (bottom_nav_menu.xml), as definições de strings (strings.xml), paleta de cores estruturada (colors.xml), estilos e temas visuais (themes.xml), drawables customizados de background e indicadores de progresso orçamentário, além do manifesto de configuração da aplicação Android (AndroidManifest.xml).")
add_code_block(doc, "app/src/main/res/menu/bottom_nav_menu.xml", code_bottom_nav_menu)
add_code_block(doc, "app/src/main/res/values/strings.xml", code_strings)
add_code_block(doc, "app/src/main/res/values/colors.xml", code_colors)
add_code_block(doc, "app/src/main/res/values/themes.xml", code_themes)

add_code_block(doc, "app/src/main/res/drawable/bg_header_blue.xml", code_bg_header_blue)
add_code_block(doc, "app/src/main/res/drawable/bg_hero_card.xml", code_bg_hero_card)
add_code_block(doc, "app/src/main/res/drawable/bg_button_primary.xml", code_bg_button_primary)
add_code_block(doc, "app/src/main/res/drawable/bg_button_secondary.xml", code_bg_button_secondary)
add_code_block(doc, "app/src/main/res/drawable/bg_button_white.xml", code_bg_button_white)
add_code_block(doc, "app/src/main/res/drawable/bg_chip_cost.xml", code_bg_chip_cost)
add_code_block(doc, "app/src/main/res/drawable/bg_chip_cost_warning.xml", code_bg_chip_cost_warning)
add_code_block(doc, "app/src/main/res/drawable/bg_chip_cost_danger.xml", code_bg_chip_cost_danger)
add_code_block(doc, "app/src/main/res/drawable/bg_edittext.xml", code_bg_edittext)
add_code_block(doc, "app/src/main/res/drawable/bg_spinner_wrapper.xml", code_bg_spinner_wrapper)
add_code_block(doc, "app/src/main/res/drawable/bg_custo_estimado.xml", code_bg_custo_estimado)
add_code_block(doc, "app/src/main/res/drawable/progress_budget.xml", code_progress_budget)
add_code_block(doc, "app/src/main/res/drawable/progress_budget_warning.xml", code_progress_budget_warning)
add_code_block(doc, "app/src/main/res/drawable/progress_budget_danger.xml", code_progress_budget_danger)

add_code_block(doc, "app/src/main/AndroidManifest.xml", code_manifest)

# ==================== 5. CONCLUSÃO ====================
doc.add_page_break()
add_heading_1(doc, "5. CONCLUSÃO")
add_body_p(doc, "O desenvolvimento do projeto FinOps Simulator viabilizou a aplicação integrada e prática dos conceitos fundamentais de Engenharia de Software, Modelagem Relacional e Programação para Dispositivos Móveis na plataforma Android. A rigorosa implementação do padrão arquitetural MVC (Model-View-Controller) demonstrou grande eficácia na obtenção de um código coeso, desacoplado e manutenível. A camada Model (com.lucasjsm.finops_simulator.model) centralizou com êxito todas as regras de domínio, validações financeiras e operações relacionais em banco SQLite; a camada View (res/layout e com.lucasjsm.finops_simulator.adapter) concentrou a apresentação visual e a experiência de uso moderna através de Material Components e ListView otimizado; e a camada Controller (com.lucasjsm.finops_simulator.activity) exerceu o controle enxuto dos fluxos operacionais e transição de telas através de Intents.")
add_body_p(doc, "No âmbito do gerenciamento de dados local, a utilização do SQLiteOpenHelper aliada ao padrão Data Access Object (DAO) demonstrou robustez técnica ao assegurar transações seguras, exclusão em cascata para garantia de integridade referencial e consultas dinâmicas filtradas por provedor de nuvem. Ademais, a incorporação de premissas reais da indústria de Cloud FinOps — como a equivalência de 730 horas operacionais por mês para serviços ininterruptos e a sinalização cromática antecipada de saturação de orçamento — conferiu à aplicação valor prático e didático expressivo, cumprindo integralmente os objetivos acadêmicos e pedagógicos propostos na disciplina.")

# ==================== 6. REFERÊNCIAS ====================
doc.add_page_break()
add_heading_1(doc, "6. REFERÊNCIAS BIBLIOGRÁFICAS")

refs = [
    "ANDROID DEVELOPERS. Guia do Desenvolvedor Android: Salvar dados usando o SQLite e Interface do Usuário. Google, 2026. Disponível em: <https://developer.android.com/guide>. Acesso em: 22 out. 2026.",
    "DEITEL, Paul; DEITEL, Harvey; DEITEL, Alexander. Android para Programadores: Uma Abordagem Baseada em Aplicativos. 3. ed. Porto Alegre: Bookman, 2016.",
    "FINOPS FOUNDATION. What is FinOps? Building Cloud Financial Management. San Francisco: The Linux Foundation, 2026. Disponível em: <https://www.finops.org/introduction/what-is-finops/>. Acesso em: 22 out. 2026.",
    "NASCIMENTO, Carlos A. Sicsú A. do. Notas de Aula e Material Didático da Disciplina de Programação para Dispositivos Móveis (Arquitetura MVC e Persistência SQLite no Android). Rio de Janeiro: Universidade do Estado do Rio de Janeiro (UERJ), Faculdade de Engenharia, 2026.",
    "SQLITE CONSORTIUM. SQLite Documentation. 2026. Disponível em: <https://www.sqlite.org/docs.html>. Acesso em: 22 out. 2026.",
    "STORMENT, J. R.; FULLER, Mike. Cloud FinOps: Collaborative, Real-Time Cloud Value Decision Making. 2. ed. Sebastopol: O'Reilly Media, 2023.",
    "[INSERIR AQUI OUTRAS REFERÊNCIAS BIBLIOGRÁFICAS UTILIZADAS PELO ALUNO, SE HOUVER]"
]

for ref in refs:
    p_ref = doc.add_paragraph()
    p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ref.paragraph_format.space_before = Pt(4)
    p_ref.paragraph_format.space_after = Pt(6)
    p_ref.paragraph_format.line_spacing = 1.0
    run_ref = p_ref.add_run(ref)
    run_ref.font.name = 'Arial'
    run_ref.font.size = Pt(11)

docx_path = BASE_DIR / "Documentacao_Academica_FinOps_Simulator_ABNT.docx"
doc.save(str(docx_path))
print(f"Arquivo DOCX gerado com sucesso: {docx_path}")

print("3/5: Construindo HTML para geração do PDF via Edge Headless...")

def escape_code(c):
    return html.escape(c)

html_parts = []
html_parts.append("""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<title>Documentação Acadêmica FinOps Simulator</title>
<style>
@page {
    size: A4 portrait;
    margin-top: 3cm;
    margin-left: 3cm;
    margin-right: 2cm;
    margin-bottom: 2cm;
}
* {
    box-sizing: border-box;
}
body {
    font-family: Arial, Helvetica, sans-serif;
    font-size: 12pt;
    line-height: 1.5;
    color: #000000;
    margin: 0;
    padding: 0;
}
.page-break {
    page-break-after: always;
    break-after: page;
}
.text-center {
    text-align: center;
}
.text-justify {
    text-align: justify;
}
.bold {
    font-weight: bold;
}
p.body-text {
    text-indent: 1.25cm;
    margin-top: 0;
    margin-bottom: 6pt;
    text-align: justify;
    line-height: 1.5;
}
p.bullet-item {
    margin-top: 0;
    margin-bottom: 4pt;
    margin-left: 1.25cm;
    text-indent: -0.5cm;
    text-align: justify;
    line-height: 1.5;
}
h1.chapter-title {
    font-size: 13pt;
    font-weight: bold;
    text-transform: uppercase;
    margin-top: 14pt;
    margin-bottom: 8pt;
    page-break-after: avoid;
    break-after: avoid;
}
h2.section-title {
    font-size: 12pt;
    font-weight: bold;
    margin-top: 14pt;
    margin-bottom: 6pt;
    page-break-after: avoid;
    break-after: avoid;
}
h3.subsection-title {
    font-size: 12pt;
    font-weight: bold;
    margin-top: 10pt;
    margin-bottom: 4pt;
    page-break-after: avoid;
    break-after: avoid;
}
p.file-label {
    font-family: Arial, sans-serif;
    font-size: 10pt;
    font-weight: bold;
    color: #1e293b;
    margin-top: 12pt;
    margin-bottom: 3pt;
    page-break-after: avoid;
    break-after: avoid;
}
pre.code-box {
    font-family: Consolas, 'Courier New', monospace;
    font-size: 8.5pt;
    line-height: 1.25;
    background-color: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 4px;
    padding: 10px;
    margin-top: 0;
    margin-bottom: 12pt;
    white-space: pre-wrap;
    word-break: break-all;
    tab-size: 4;
}
.screenshot-box {
    border: 1.5px dashed #94a3b8;
    background-color: #f1f5f9;
    padding: 24px;
    text-align: center;
    font-weight: bold;
    font-style: italic;
    color: #334155;
    margin: 16pt 0;
    border-radius: 6px;
    font-size: 11pt;
    page-break-inside: avoid;
}
.natureza-box {
    margin-left: 8cm;
    font-size: 10pt;
    line-height: 1.15;
    text-align: justify;
    margin-top: 40pt;
    margin-bottom: 140pt;
}
.ref-item {
    font-size: 11pt;
    line-height: 1.2;
    text-align: justify;
    margin-bottom: 8pt;
}
.sumario-item {
    font-size: 11pt;
    line-height: 1.35;
    margin-bottom: 2pt;
}
</style>
</head>
<body>
""")

# Capa
html_parts.append("""
<div class="text-center" style="padding-top: 10pt;">
    <div class="bold" style="font-size: 12pt;">UNIVERSIDADE DO ESTADO DO RIO DE JANEIRO — UERJ</div>
    <div class="bold" style="font-size: 12pt;">FACULDADE DE ENGENHARIA — FEN</div>
    <div class="bold" style="font-size: 12pt; margin-bottom: 120pt;">DEPARTAMENTO DE ENGENHARIA DE SISTEMAS E COMPUTAÇÃO</div>
    
    <div class="bold" style="font-size: 12pt; margin-bottom: 160pt;">[INSERIR NOME COMPLETO DO ALUNO — MATRÍCULA: XXXXXXXXX]</div>
    
    <div class="bold" style="font-size: 14pt; margin-bottom: 200pt; line-height: 1.4;">
        FINOPS SIMULATOR: SISTEMA MÓVEL DE SIMULAÇÃO E GESTÃO DE CUSTOS DE NUVEM COMPUTACIONAL
    </div>
    
    <div style="font-size: 12pt;">Rio de Janeiro</div>
    <div style="font-size: 12pt;">2026</div>
</div>
<div class="page-break"></div>
""")

# Folha de Rosto
html_parts.append("""
<div class="text-center" style="padding-top: 10pt;">
    <div class="bold" style="font-size: 12pt; margin-bottom: 60pt;">FOLHA DE ROSTO (CONTRACAPA)</div>
    <div class="bold" style="font-size: 12pt; margin-bottom: 100pt;">[INSERIR NOME COMPLETO DO ALUNO — MATRÍCULA: XXXXXXXXX]</div>
    
    <div class="bold" style="font-size: 14pt; margin-bottom: 40pt; line-height: 1.4;">
        FINOPS SIMULATOR: SISTEMA MÓVEL DE SIMULAÇÃO E GESTÃO DE CUSTOS DE NUVEM COMPUTACIONAL
    </div>
    
    <div class="natureza-box">
        Trabalho acadêmico apresentado à disciplina de Programação para Dispositivos Móveis, do Curso de Engenharia de Sistemas e Computação da Universidade do Estado do Rio de Janeiro (UERJ), como requisito parcial de avaliação.
        <br><br>
        <strong>Professor Orientador: Prof. Dr. Engº Carlos A. Sicsú A. do Nascimento</strong>
    </div>
    
    <div style="font-size: 12pt;">Rio de Janeiro</div>
    <div style="font-size: 12pt;">2026</div>
</div>
<div class="page-break"></div>
""")

# Sumário
html_parts.append("""
<div>
    <div class="text-center bold" style="font-size: 14pt; margin-bottom: 24pt;">SUMÁRIO (ÍNDICE)</div>
""")
for texto, bold, indent in sumario_linhas:
    b_start = "<strong>" if bold else ""
    b_end = "</strong>" if bold else ""
    ml = f"margin-left: {indent*2}cm;" if indent > 0 else ""
    html_parts.append(f'<div class="sumario-item" style="{ml}">{b_start}{texto}{b_end}</div>\n')

html_parts.append("""
</div>
<div class="page-break"></div>
""")

# 1. Introdução
html_parts.append("""
<h1 class="chapter-title">1. INTRODUÇÃO</h1>

<h2 class="section-title">1.1 Contextualização e Tema do Projeto</h2>
<p class="body-text">A crescente adoção do paradigma de computação em nuvem (Cloud Computing) transformou profundamente a arquitetura, o provisionamento e a operação de sistemas de software corporativos modernos. A migração de ambientes locais físicos (on-premises) para os principais provedores de nuvem pública do mercado global — notadamente Amazon Web Services (AWS), Google Cloud Platform (GCP), Microsoft Azure e Oracle Cloud Infrastructure (OCI) — proporcionou às organizações elasticidade operacional incomparável, alta disponibilidade regional e capacidade de rápido escalonamento de aplicações. Todavia, a transição para um modelo financeiro puramente variável baseado no consumo efetivo sob demanda (pay-as-you-go) trouxe consigo novos e severos desafios de previsibilidade orçamentária, tornando comum o surgimento de gastos imprevistos e estouros recorrentes de orçamento (budget overruns).</p>
<p class="body-text">Nesse contexto de complexidade operacional e financeira, consolidou-se internacionalmente a disciplina de Cloud FinOps (Financial Operations em Nuvem), padronizada e fomentada pela FinOps Foundation. O FinOps representa uma evolução cultural e técnica interdisciplinar que combina práticas financeiras tradicionais, engenharia de confiabilidade e operações DevOps, capacitando equipes de desenvolvimento e engenharia a assumirem a responsabilidade financeira direta pelos recursos computacionais que provisionam. Seu objetivo fundamental não é apenas conter custos de maneira indiscriminada, mas sim viabilizar decisões arquiteturais fundamentadas em valor de negócio, garantindo máxima eficiência por unidade monetária investida em infraestrutura.</p>
<p class="body-text">Com o propósito de aplicar e consolidar esses conceitos em um ambiente prático na plataforma móvel Android, foi concebido e implementado o aplicativo FinOps Simulator (desenvolvido integralmente em linguagem Java). O sistema atua como um simulador estruturado de custos e ferramenta móvel de governança orçamentária multicloud, voltado a engenheiros de software, arquitetos de nuvem, gestores de tecnologia e estudantes universitários.</p>
<p class="body-text">O aplicativo permite cadastrar e manter projetos independentes associados às quatro maiores plataformas de nuvem do mercado (AWS, GCP, Azure e OCI), com estipulação de teto orçamentário mensal (budget). A partir de um catálogo pré-carregado contendo 14 serviços emblemáticos de infraestrutura (como instâncias elásticas de computação, armazenamento de objetos, data warehouses analíticos de alto desempenho, bancos de dados NoSQL e clusters Kubernetes gerenciados), o usuário pode dimensionar e alocar recursos a cada projeto. O simulador apura de maneira instantânea o custo financeiro estimado com base nas métricas reais de cobrança (incluindo o fator temporal de 730 horas mensais para serviços contínuos), apresentando em tempo real o percentual de consumo do orçamento e sinalizando preventivamente alertas visuais cromáticos (faixas de conformidade, atenção a partir de 80% e perigo/estouro acima de 100%).</p>

<h2 class="section-title">1.2 Objetivos Geral e Específicos</h2>
<p class="body-text">O objetivo geral deste trabalho acadêmico consiste em projetar, arquitetar e implementar uma aplicação Android nativa e completa, desenvolvida em linguagem Java, adotando o padrão arquitetural MVC (Model-View-Controller) e utilizando persistência local em banco de dados relacional SQLite, consolidando na prática as competências e metodologias fundamentais de desenvolvimento para dispositivos móveis ministradas na disciplina.</p>
<p class="body-text">Como objetivos específicos que nortearam o desenvolvimento do projeto, destacam-se:</p>
<p class="bullet-item">- Implementar operações completas de CRUD (Create, Read, Update, Delete) para as três entidades nucleares do sistema: Projeto, Tipo de Recurso (Catálogo) e Recurso Alocado (tabela associativa);</p>
<p class="bullet-item">- Isolar com rigor as regras de negócio, a lógica matemática de precificação em nuvem (fator de 730 horas/mês, conversões de unidade e somatórios agregados) e as manipulações relacionais SQL na camada Model (com.lucasjsm.finops_simulator.model);</p>
<p class="bullet-item">- Construir interfaces gráficas responsivas, intuitivas e aderentes ao Material Design na camada View (res/layout e com.lucasjsm.finops_simulator.adapter), utilizando ListView associado ao padrão de reciclagem ViewHolder, MaterialCardView, ProgressBar dinâmico, ScrollView, Spinner e BottomNavigationView;</p>
<p class="bullet-item">- Implementar filtros combinados e reativos no Catálogo de Serviços, cálculo instantâneo em tempo real de custo (live preview) nos formulários de cadastro através do padrão TextWatcher, e agregação dinâmica automática de custos no painel de controle;</p>
<p class="bullet-item">- Assegurar a integridade relacional e a consistência histórica do banco de dados SQLite, implementando exclusão em cascata (cascade delete) para eliminar automaticamente os recursos vinculados ao remover um projeto, além de salvaguardar regras de domínio como a imutabilidade do provedor do projeto e o bloqueio de serviço na alteração de alocações.</p>

<h2 class="section-title">1.3 Padrão Arquitetural MVC e Persistência SQLite</h2>
<p class="body-text">Em rigorosa conformidade com o programa curricular e a metodologia pedagógica apresentada em aula pelo Prof. Dr. Engº Carlos A. Sicsú A. do Nascimento, o FinOps Simulator adota a separação clássica de responsabilidades em três camadas fundamentais:</p>
<p class="bullet-item">- <strong>Model (com.lucasjsm.finops_simulator.model):</strong> Contém as classes de entidade de domínio (Projeto, TipoRecurso e RecursoAlocado), a classe centralizadora de gerenciamento do banco de dados SQLite (DatabaseHelper, subclasse de SQLiteOpenHelper) e as classes especializadas de acesso e persistência a dados (ProjetoDAO, TipoRecursoDAO e RecursoAlocadoDAO). Toda a lógica financeira de cálculo de custo de nuvem, conversão de horas operacionais, somatórios agregados, consultas parametrizadas com filtros e operações transacionais com exclusão em cascata residem exclusivamente nesta camada, sem qualquer dependência das bibliotecas visuais do Android.</p>
<p class="bullet-item">- <strong>View (res/layout/*.xml e com.lucasjsm.finops_simulator.adapter):</strong> Compreende a totalidade dos arquivos XML declarativos de layout, drawables vetoriais customizados, paleta semântica de cores, temas visuais e as quatro classes de adaptadores de lista (HomeProjetoAdapter, ProjetoAdapter, RecursoAdapter e TipoRecursoAdapter). Os adaptadores estendem ArrayAdapter e utilizam o padrão estático ViewHolder para evitar chamadas redundantes a findViewById(), sendo responsáveis estritamente pela renderização dos dados fornecidos pelo Model, estilização dinâmica dos indicadores visuais de orçamento e disparo de eventos de interface (através da interface OnRecursoActionListener e de cliques convencionais).</p>
<p class="bullet-item">- <strong>Controller (com.lucasjsm.finops_simulator.activity):</strong> Composto pelas seis Activities do aplicativo (MainActivity, ListaProjetosActivity, FormProjetoActivity, DashboardActivity, FormRecursoActivity e ListaTiposActivity). As Activities atuam exclusivamente como mediadoras e coordenadoras de fluxo: capturam os eventos e interações do usuário disparados na View, coordenam o ciclo de vida do Android (onCreate, onResume, onPause), recebem e repassam dados via Intent (utilizando parâmetros padronizados como acao = -1 para inclusão e acao = 0 para alteração), invocam as rotinas de persistência e consulta dos DAOs e atualizam a View com os dados retornados.</p>
<div class="page-break"></div>
""")

# 2. Arquitetura
html_parts.append("""
<h1 class="chapter-title">2. ARQUITETURA DO SISTEMA E MODELAGEM DE DADOS</h1>

<h2 class="section-title">2.1 Divisão em Pacotes segundo o Padrão MVC</h2>
<p class="body-text">A estrutura de código-fonte do projeto FinOps Simulator está organizada de maneira modularizada e coesa nos seguintes pacotes Java:</p>
<p class="bullet-item">- <strong>com.lucasjsm.finops_simulator.model:</strong> DatabaseHelper, Projeto, ProjetoDAO, TipoRecurso, TipoRecursoDAO, RecursoAlocado e RecursoAlocadoDAO;</p>
<p class="bullet-item">- <strong>com.lucasjsm.finops_simulator.adapter:</strong> HomeProjetoAdapter, ProjetoAdapter, RecursoAdapter e TipoRecursoAdapter, complementadas pelas definições visuais em res/layout/ e drawables em res/drawable/;</p>
<p class="bullet-item">- <strong>com.lucasjsm.finops_simulator.activity:</strong> MainActivity, ListaProjetosActivity, FormProjetoActivity, DashboardActivity, FormRecursoActivity e ListaTiposActivity.</p>

<h2 class="section-title">2.2 Esquema Relacional do Banco de Dados SQLite (finops_simulator.db)</h2>
<p class="body-text">O banco de dados local da aplicação, denominado finops_simulator.db (versão 1), é inteiramente administrado pela classe DatabaseHelper. O esquema relacional é estruturado em três tabelas associadas por integridade referencial:</p>
<p class="bullet-item">1. <strong>Tabela projeto:</strong> Armazena os dados cadastrais e orçamentários dos projetos simulados. Possui os campos id (INTEGER PRIMARY KEY AUTOINCREMENT), nome (TEXT NOT NULL), descricao (TEXT), provedor (TEXT NOT NULL), orcamento_maximo (REAL) e data_criacao (TEXT). Apresenta relacionamento de cardinalidade 1 para N com a tabela recurso_alocado.</p>
<p class="bullet-item">2. <strong>Tabela tipo_recurso:</strong> Representa o catálogo pré-cadastrado de serviços de nuvem computacional. Possui as colunas id (INTEGER PRIMARY KEY AUTOINCREMENT), nome (TEXT NOT NULL), provedor (TEXT NOT NULL), unidade (TEXT NOT NULL), custo_unitario (REAL NOT NULL) e descricao (TEXT). No momento da criação do banco, são automaticamente injetados 14 registros oficiais de serviços das quatro maiores plataformas globais (AWS, GCP, Azure e OCI), cobrindo categorias de computação, armazenamento, bancos de dados e analytics.</p>
<p class="bullet-item">3. <strong>Tabela recurso_alocado:</strong> Tabela associativa que relaciona um projeto a um tipo de recurso dimensionado, viabilizando o relacionamento relacional N para N entre as entidades através de uma entidade associativa com atributos próprios. Contém os campos id (INTEGER PRIMARY KEY AUTOINCREMENT), projeto_id (INTEGER NOT NULL, chave estrangeira para projeto.id), tipo_recurso_id (INTEGER NOT NULL, chave estrangeira para tipo_recurso.id), quantidade (REAL NOT NULL), periodo (REAL NOT NULL), custo_estimado (REAL) e observacao (TEXT).</p>

<h2 class="section-title">2.3 Regras de Negócio e Integridade de Domínio</h2>
<p class="body-text">A integridade lógica do FinOps Simulator é assegurada por regras de negócio rigorosas e proteções de domínio incorporadas na camada Model:</p>
<p class="bullet-item">- <strong>Cálculo de Custo Estimado com Fator de Equivalência Temporal (730 h/mês):</strong> No método calcularCusto(custoUnitario, unidadeRecurso) da classe RecursoAlocado, o sistema implementa a convenção oficial da indústria de computação em nuvem para serviços cobrados por hora contínua. Como o ano civil médio possui 365 dias, divididos por 12 meses e multiplicados por 24 horas/dia, obtém-se o fator padronizado de 730 horas operacionais por mês. Caso a unidade seja 'hora', o multiplicador de tempo adotado é 730.0; para grandezas pontuais (como GB, TB ou milhões de requisições), o multiplicador é 1.0. A fórmula implementada é: custoEstimado = quantidade * custoUnitario * multiplicadorTempo * periodo.</p>
<p class="bullet-item">- <strong>Imutabilidade do Provedor de Nuvem no Projeto:</strong> Ao criar um projeto, o provedor (AWS, GCP, Azure ou OCI) é fixado. Caso o usuário pudesse alterar o provedor de um projeto em andamento, os recursos de infraestrutura já alocados (ex.: instâncias AWS EC2) passariam a constar sob outro provedor (ex.: Google Cloud), gerando uma inconsistência de domínio inaceitável. Dessa forma, na tela de edição (FormProjetoActivity), o componente Spinner de seleção de provedor é desabilitado (spProvedor.setEnabled(false)), e o método ProjetoDAO.alterar() atualiza exclusivamente os campos nome, descricao e orcamento_maximo.</p>
<p class="bullet-item">- <strong>Isolamento Contextual de Serviços por Provedor:</strong> Ao alocar um recurso no projeto (FormRecursoActivity), o catálogo de serviços apresentado para seleção é filtrado de maneira estrita pelo provedor do projeto através de tipoDao.getAllPorProvedor(projeto.getProvedor()), impedindo que serviços da AWS sejam indevidamente alocados em projetos da Azure ou OCI.</p>
<p class="bullet-item">- <strong>Bloqueio do Tipo de Recurso na Edição de Alocação:</strong> Quando uma alocação de recurso já existente é editada, o tipo de recurso associado permanece fixo (spTipo.setEnabled(false)), permitindo ao usuário alterar o volume, o período operacional em meses e as observações textuais, disparando o recálculo automático do custo estimado com garantia de rastreabilidade.</p>
<p class="bullet-item">- <strong>Exclusão em Cascata (Cascade Delete) no SQLite:</strong> Ao excluir um projeto, o método ProjetoDAO.apagar(long id) executa a deleção prévia de todos os registros na tabela recurso_alocado cujo projeto_id coincida com o identificador do projeto antes de remover o próprio registro da tabela projeto, impedindo a existência de registros órfãos e mantendo a sanidade relacional do banco de dados local.</p>
<p class="bullet-item">- <strong>Cálculo Consolidado e Prevenção Visual de Estouro Orçamentário:</strong> O método recursoDao.calcularCustoTotal(long projetoId) recupera todos os recursos do projeto e computa o somatório de custoEstimado. Com base na proporção entre o custo total e o orçamento estipulado (pct = (custo / orcamento) * 100), as telas aplicam estilos cromáticos reativos de sinalização: Normal (&lt; 80%, cores primárias e tons de verde saudável), Alerta (80% a 100%, tons em amarelo/âmbar com texto #B45309) e Estouro/Perigo (&gt; 100%, tons em carmim/vermelho com texto #BE123C), alertando o usuário em tempo real sobre a necessidade de readequação de gastos.</p>
<div class="page-break"></div>
""")

# 3. Model
html_parts.append("""
<h1 class="chapter-title">3. CAMADA MODEL — CLASSES DE ENTIDADE E PERSISTÊNCIA (DAO)</h1>
<p class="body-text">A seguir são apresentados todos os códigos-fonte Java pertencentes ao pacote com.lucasjsm.finops_simulator.model, responsáveis pela integridade do esquema SQLite, representação das entidades de domínio e operações de persistência e cálculo.</p>

<h2 class="section-title">3.1 Gerenciador do Banco de Dados (DatabaseHelper.java)</h2>
<p class="body-text">Classe responsável pela instanciação do banco de dados finops_simulator.db, criação das tabelas projeto, tipo_recurso e recurso_alocado com suas respectivas restrições de integridade referencial, e injeção do catálogo inicial com 14 serviços de nuvem computacional.</p>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/model/DatabaseHelper.java</p>
<pre class="code-box"><code>""" + escape_code(code_DatabaseHelper) + """</code></pre>

<h2 class="section-title">3.2 Entidade e Persistência de Projeto (Projeto.java e ProjetoDAO.java)</h2>
<p class="body-text">A classe de entidade Projeto encapsula os atributos de um projeto de nuvem (nome, descrição, provedor, teto orçamentário mensal e data de criação), além do custo consolidado calculado. A classe ProjetoDAO concentra as operações de CRUD, consultas parametrizadas ordenadas por nome, atualização com proteção de imutabilidade do provedor e rotina de exclusão em cascata.</p>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/model/Projeto.java</p>
<pre class="code-box"><code>""" + escape_code(code_Projeto) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/model/ProjetoDAO.java</p>
<pre class="code-box"><code>""" + escape_code(code_ProjetoDAO) + """</code></pre>

<h2 class="section-title">3.3 Entidade e Persistência de Tipo de Recurso (TipoRecurso.java e TipoRecursoDAO.java)</h2>
<p class="body-text">A classe TipoRecurso modela as especificações de um serviço do catálogo de nuvem (nome, provedor, unidade de medida, custo unitário e descrição técnica). A classe TipoRecursoDAO realiza as consultas ao catálogo, incluindo a listagem geral ordenada e a consulta parametrizada com cláusula WHERE filtrada estritamente pelo provedor de nuvem informado.</p>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/model/TipoRecurso.java</p>
<pre class="code-box"><code>""" + escape_code(code_TipoRecurso) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/model/TipoRecursoDAO.java</p>
<pre class="code-box"><code>""" + escape_code(code_TipoRecursoDAO) + """</code></pre>

<h2 class="section-title">3.4 Entidade e Persistência de Recurso Alocado (RecursoAlocado.java e RecursoAlocadoDAO.java)</h2>
<p class="body-text">A classe RecursoAlocado modela a alocação de um serviço específico a um projeto, encapsulando a regra de negócio do cálculo de custo estimado baseado na métrica da unidade (aplicando o fator de 730 horas/mês para serviços contínuos). A classe RecursoAlocadoDAO implementa a persistência da tabela associativa, listagem de recursos vinculados a um projeto e cálculo agregado do custo total do projeto.</p>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/model/RecursoAlocado.java</p>
<pre class="code-box"><code>""" + escape_code(code_RecursoAlocado) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/model/RecursoAlocadoDAO.java</p>
<pre class="code-box"><code>""" + escape_code(code_RecursoAlocadoDAO) + """</code></pre>
<div class="page-break"></div>
""")

# 4. Apresentação das Telas
html_parts.append("""
<h1 class="chapter-title">4. APRESENTAÇÃO DAS TELAS E CÓDIGOS-FONTE (VIEW E CONTROLLER)</h1>
<p class="body-text">Neste capítulo são detalhadas todas as telas da aplicação FinOps Simulator. Para cada interface gráfica, apresentam-se a descrição funcional completa, o espaço reservado para a inserção da captura de tela (print) e a totalidade dos códigos-fonte associados em XML (camada View) e Java (camadas Controller/Activity e View/Adapter).</p>
<div class="page-break"></div>

<h2 class="section-title">4.1 Tela Principal — Visão Geral e Acesso Rápido (MainActivity)</h2>
<p class="body-text">Descrição Funcional: Tela inicial do aplicativo FinOps Simulator. Apresenta o cabeçalho institucional do sistema com ícone temático de nuvem, um Hero Card estilizado em gradiente azul com atalho rápido para criação de novos projetos, uma seção informativa de seleção e a listagem dinâmica de projetos cadastrados em um ListView. Cada item da lista exibe o número sequencial do projeto, nome, provedor de nuvem, orçamento mensal configurado, chip cromático com o custo total mensal e o percentual de consumo do orçamento (com formatação condicional em verde, amarelo para &gt;= 80% e vermelho para &gt; 100%), além de dois botões de ação rápida: 'Abrir Dashboard' e '+ Recursos'. A navegação inferior é provida pelo BottomNavigationView padronizado.</p>
<div class="screenshot-box">[INSERIR AQUI O PRINT DA TELA 1: TELA PRINCIPAL — VISÃO GERAL (MainActivity)]</div>
<p class="file-label">Arquivo: app/src/main/res/layout/activity_main.xml</p>
<pre class="code-box"><code>""" + escape_code(code_activity_main) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/layout/header_home.xml</p>
<pre class="code-box"><code>""" + escape_code(code_header_home) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/layout/item_projeto_home.xml</p>
<pre class="code-box"><code>""" + escape_code(code_item_projeto_home) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/activity/MainActivity.java</p>
<pre class="code-box"><code>""" + escape_code(code_MainActivity) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/adapter/HomeProjetoAdapter.java</p>
<pre class="code-box"><code>""" + escape_code(code_HomeProjetoAdapter) + """</code></pre>
<div class="page-break"></div>

<h2 class="section-title">4.2 Tela de Lista de Projetos (ListaProjetosActivity)</h2>
<p class="body-text">Descrição Funcional: Tela dedicada à gestão e visualização abrangente de todos os projetos cadastrados no simulador. Possui cabeçalho superior estilizado em azul com botão circular de retorno, subtítulo descritivo, ListView com divisor transparente e espaçamento harmônico entre cards, estado vazio amigável quando não há projetos registrados e botão de ação flutuante (FAB) para cadastro de novo projeto. Cada card exibe o nome do projeto, tag do provedor, custo estimado total mensal, teto orçamentário estipulado e uma barra de progresso horizontal (ProgressBar) com preenchimento cromático dinâmico que reflete o percentual de consumo do orçamento.</p>
<div class="screenshot-box">[INSERIR AQUI O PRINT DA TELA 2: TELA DE LISTA DE PROJETOS (ListaProjetosActivity)]</div>
<p class="file-label">Arquivo: app/src/main/res/layout/activity_lista_projetos.xml</p>
<pre class="code-box"><code>""" + escape_code(code_activity_lista_projetos) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/layout/item_projeto.xml</p>
<pre class="code-box"><code>""" + escape_code(code_item_projeto) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/activity/ListaProjetosActivity.java</p>
<pre class="code-box"><code>""" + escape_code(code_ListaProjetosActivity) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/adapter/ProjetoAdapter.java</p>
<pre class="code-box"><code>""" + escape_code(code_ProjetoAdapter) + """</code></pre>
<div class="page-break"></div>

<h2 class="section-title">4.3 Tela de Cadastro e Edição de Projeto (FormProjetoActivity)</h2>
<p class="body-text">Descrição Funcional: Formulário versátil utilizado tanto para a criação de um novo projeto (quando recebe acao = -1 via Intent) quanto para a edição de um projeto existente (quando recebe acao = 0 e o identificador do projeto). Contém campos para Nome do Projeto, Descrição, seletor de Provedor de Nuvem via Spinner (AWS, GCP, Azure, OCI) e campo numérico decimal para estipulação do Orçamento Máximo mensal em reais. Na modalidade de edição, o seletor de provedor é automaticamente bloqueado (spProvedor.setEnabled(false)) para resguardar a integridade dos serviços já alocados, e o botão de persistência aciona a rotina de atualização correspondente no ProjetoDAO.</p>
<div class="screenshot-box">[INSERIR AQUI O PRINT DA TELA 3: TELA DE CADASTRO E EDIÇÃO DE PROJETO (FormProjetoActivity)]</div>
<p class="file-label">Arquivo: app/src/main/res/layout/activity_form_projeto.xml</p>
<pre class="code-box"><code>""" + escape_code(code_activity_form_projeto) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/activity/FormProjetoActivity.java</p>
<pre class="code-box"><code>""" + escape_code(code_FormProjetoActivity) + """</code></pre>
<div class="page-break"></div>

<h2 class="section-title">4.4 Tela de Dashboard e Gestão do Projeto (DashboardActivity)</h2>
<p class="body-text">Descrição Funcional: Painel analítico de controle de um projeto específico. Apresenta cabeçalho estilizado, card de síntese financeira com nome, provedor, teto orçamentário, custo total consumido e barra de progresso reativa com rótulo percentual. Na seção inferior, lista os recursos alocados àquele projeto através de um ListView customizado, com indicação de serviço, volume dimensionado, período em meses, custo parcial estimado e botão de exclusão individual. Na extremidade inferior direita, disponibiliza botões de ação empilhados para exclusão integral do projeto (com diálogo de confirmação AlertDialog e exclusão em cascata de recursos), edição dos dados do projeto e adição de novos recursos computacionais.</p>
<div class="screenshot-box">[INSERIR AQUI O PRINT DA TELA 4: TELA DE DASHBOARD E GESTÃO DO PROJETO (DashboardActivity)]</div>
<p class="file-label">Arquivo: app/src/main/res/layout/activity_dashboard.xml</p>
<pre class="code-box"><code>""" + escape_code(code_activity_dashboard) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/layout/item_recurso.xml</p>
<pre class="code-box"><code>""" + escape_code(code_item_recurso) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/activity/DashboardActivity.java</p>
<pre class="code-box"><code>""" + escape_code(code_DashboardActivity) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/adapter/RecursoAdapter.java</p>
<pre class="code-box"><code>""" + escape_code(code_RecursoAdapter) + """</code></pre>
<div class="page-break"></div>

<h2 class="section-title">4.5 Tela de Alocação e Edição de Recurso (FormRecursoActivity)</h2>
<p class="body-text">Descrição Funcional: Formulário responsável pelo dimensionamento de recursos de infraestrutura vinculados a um projeto. Suporta os modos de inclusão (acao = -1) e alteração (acao = 0). O seletor Spinner de tipo de recurso carrega exclusivamente os serviços pertinentes ao provedor do projeto em questão. Apresenta campos de entrada para Quantidade/Volume (em GB, instâncias ou TB), Período operacional em meses e campo opcional para Observações. Através de ouvintes TextWatcher associados aos campos de digitação e do ouvinte de seleção do Spinner, a tela atualiza em tempo real uma caixa de pré-visualização (live preview) do custo estimado resultante da alocação, aplicando o fator de 730 horas/mês quando o serviço for de tarifação horária.</p>
<div class="screenshot-box">[INSERIR AQUI O PRINT DA TELA 5: TELA DE ALOCAÇÃO E EDIÇÃO DE RECURSO (FormRecursoActivity)]</div>
<p class="file-label">Arquivo: app/src/main/res/layout/activity_form_recurso.xml</p>
<pre class="code-box"><code>""" + escape_code(code_activity_form_recurso) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/activity/FormRecursoActivity.java</p>
<pre class="code-box"><code>""" + escape_code(code_FormRecursoActivity) + """</code></pre>
<div class="page-break"></div>

<h2 class="section-title">4.6 Tela de Catálogo de Serviços de Nuvem (ListaTiposActivity)</h2>
<p class="body-text">Descrição Funcional: Catálogo de referência dos serviços de computação em nuvem pré-configurados no banco de dados. Apresenta cabeçalho informativo, seletor Spinner de filtragem com as opções 'Todos os Provedores', 'AWS', 'GCP', 'Azure' e 'OCI', e uma listagem em ListView contendo os cards individuais de serviços. Cada card apresenta o nome do serviço, chip com coloração temática correspondente ao provedor (âmbar para AWS, azul royal para GCP, azul celeste para Azure e carmim para OCI), custo unitário formatado com precisão de três casas decimais por unidade métrica e descrição técnica de finalidade. Ao tocar em um item da lista, exibe uma mensagem Toast informativa com a especificação consolidada do serviço.</p>
<div class="screenshot-box">[INSERIR AQUI O PRINT DA TELA 6: TELA DE CATÁLOGO DE SERVIÇOS DE NUVEM (ListaTiposActivity)]</div>
<p class="file-label">Arquivo: app/src/main/res/layout/activity_lista_tipos.xml</p>
<pre class="code-box"><code>""" + escape_code(code_activity_lista_tipos) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/layout/item_tipo_recurso.xml</p>
<pre class="code-box"><code>""" + escape_code(code_item_tipo_recurso) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/activity/ListaTiposActivity.java</p>
<pre class="code-box"><code>""" + escape_code(code_ListaTiposActivity) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/java/com/lucasjsm/finops_simulator/adapter/TipoRecursoAdapter.java</p>
<pre class="code-box"><code>""" + escape_code(code_TipoRecursoAdapter) + """</code></pre>
<div class="page-break"></div>

<h2 class="section-title">4.7 Utilitários de Interface, Recursos XML e Manifesto</h2>
<p class="body-text">Abaixo encontram-se o menu de navegação inferior (bottom_nav_menu.xml), as definições de strings (strings.xml), paleta de cores estruturada (colors.xml), estilos e temas visuais (themes.xml), drawables customizados de background e indicadores de progresso orçamentário, além do manifesto de configuração da aplicação Android (AndroidManifest.xml).</p>
<p class="file-label">Arquivo: app/src/main/res/menu/bottom_nav_menu.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bottom_nav_menu) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/values/strings.xml</p>
<pre class="code-box"><code>""" + escape_code(code_strings) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/values/colors.xml</p>
<pre class="code-box"><code>""" + escape_code(code_colors) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/values/themes.xml</p>
<pre class="code-box"><code>""" + escape_code(code_themes) + """</code></pre>

<p class="file-label">Arquivo: app/src/main/res/drawable/bg_header_blue.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_header_blue) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_hero_card.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_hero_card) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_button_primary.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_button_primary) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_button_secondary.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_button_secondary) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_button_white.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_button_white) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_chip_cost.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_chip_cost) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_chip_cost_warning.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_chip_cost_warning) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_chip_cost_danger.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_chip_cost_danger) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_edittext.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_edittext) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_spinner_wrapper.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_spinner_wrapper) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/bg_custo_estimado.xml</p>
<pre class="code-box"><code>""" + escape_code(code_bg_custo_estimado) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/progress_budget.xml</p>
<pre class="code-box"><code>""" + escape_code(code_progress_budget) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/progress_budget_warning.xml</p>
<pre class="code-box"><code>""" + escape_code(code_progress_budget_warning) + """</code></pre>
<p class="file-label">Arquivo: app/src/main/res/drawable/progress_budget_danger.xml</p>
<pre class="code-box"><code>""" + escape_code(code_progress_budget_danger) + """</code></pre>

<p class="file-label">Arquivo: app/src/main/AndroidManifest.xml</p>
<pre class="code-box"><code>""" + escape_code(code_manifest) + """</code></pre>
<div class="page-break"></div>
""")

# 5. Conclusão
html_parts.append("""
<h1 class="chapter-title">5. CONCLUSÃO</h1>
<p class="body-text">O desenvolvimento do projeto FinOps Simulator viabilizou a aplicação integrada e prática dos conceitos fundamentais de Engenharia de Software, Modelagem Relacional e Programação para Dispositivos Móveis na plataforma Android. A rigorosa implementação do padrão arquitetural MVC (Model-View-Controller) demonstrou grande eficácia na obtenção de um código coeso, desacoplado e manutenível. A camada Model (com.lucasjsm.finops_simulator.model) centralizou com êxito todas as regras de domínio, validações financeiras e operações relacionais em banco SQLite; a camada View (res/layout e com.lucasjsm.finops_simulator.adapter) concentrou a apresentação visual e a experiência de uso moderna através de Material Components e ListView otimizado; e a camada Controller (com.lucasjsm.finops_simulator.activity) exerceu o controle enxuto dos fluxos operacionais e transição de telas através de Intents.</p>
<p class="body-text">No âmbito do gerenciamento de dados local, a utilização do SQLiteOpenHelper aliada ao padrão Data Access Object (DAO) demonstrou robustez técnica ao assegurar transações seguras, exclusão em cascata para garantia de integridade referencial e consultas dinâmicas filtradas por provedor de nuvem. Ademais, a incorporação de premissas reais da indústria de Cloud FinOps — como a equivalência de 730 horas operacionais por mês para serviços ininterruptos e a sinalização cromática antecipada de saturação de orçamento — conferiu à aplicação valor prático e didático expressivo, cumprindo integralmente os objetivos acadêmicos e pedagógicos propostos na disciplina.</p>
<div class="page-break"></div>
""")

# 6. Referências
html_parts.append("""
<h1 class="chapter-title">6. REFERÊNCIAS BIBLIOGRÁFICAS</h1>
<div class="ref-item">ANDROID DEVELOPERS. Guia do Desenvolvedor Android: Salvar dados usando o SQLite e Interface do Usuário. Google, 2026. Disponível em: &lt;https://developer.android.com/guide&gt;. Acesso em: 22 out. 2026.</div>
<div class="ref-item">DEITEL, Paul; DEITEL, Harvey; DEITEL, Alexander. Android para Programadores: Uma Abordagem Baseada em Aplicativos. 3. ed. Porto Alegre: Bookman, 2016.</div>
<div class="ref-item">FINOPS FOUNDATION. What is FinOps? Building Cloud Financial Management. San Francisco: The Linux Foundation, 2026. Disponível em: &lt;https://www.finops.org/introduction/what-is-finops/&gt;. Acesso em: 22 out. 2026.</div>
<div class="ref-item">NASCIMENTO, Carlos A. Sicsú A. do. Notas de Aula e Material Didático da Disciplina de Programação para Dispositivos Móveis (Arquitetura MVC e Persistência SQLite no Android). Rio de Janeiro: Universidade do Estado do Rio de Janeiro (UERJ), Faculdade de Engenharia, 2026.</div>
<div class="ref-item">SQLITE CONSORTIUM. SQLite Documentation. 2026. Disponível em: &lt;https://www.sqlite.org/docs.html&gt;. Acesso em: 22 out. 2026.</div>
<div class="ref-item">STORMENT, J. R.; FULLER, Mike. Cloud FinOps: Collaborative, Real-Time Cloud Value Decision Making. 2. ed. Sebastopol: O'Reilly Media, 2023.</div>
<div class="ref-item">[INSERIR AQUI OUTRAS REFERÊNCIAS BIBLIOGRÁFICAS UTILIZADAS PELO ALUNO, SE HOUVER]</div>
""")

html_parts.append("""
</body>
</html>
""")

html_path = BASE_DIR / "Documentacao_Academica_FinOps_Simulator_ABNT.html"
with open(html_path, "w", encoding="utf-8") as f:
    f.write("".join(html_parts))

print(f"HTML gerado com sucesso: {html_path}")

print("4/5: Convertendo HTML para PDF bruto com Edge Headless...")
raw_pdf_path = BASE_DIR / "Documentacao_Academica_FinOps_Simulator_raw.pdf"

cmd = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    "--headless=new",
    "--disable-gpu",
    "--no-pdf-header-footer",
    f"--print-to-pdf={raw_pdf_path}",
    str(html_path)
]
subprocess.run(cmd, check=True)
print(f"PDF bruto gerado: {raw_pdf_path} ({raw_pdf_path.stat().st_size} bytes)")

print("5/5: Aplicando numeração de páginas ABNT no PDF...")
final_pdf_path = BASE_DIR / "Documentacao_Academica_FinOps_Simulator_ABNT.pdf"

reader = pypdf.PdfReader(str(raw_pdf_path))
writer = pypdf.PdfWriter()
total_pages = len(reader.pages)
print(f"Total de páginas no documento: {total_pages}")

for idx, page in enumerate(reader.pages):
    page_num = idx + 1
    if page_num <= 3:
        # Páginas pré-textuais sem número visível
        writer.add_page(page)
    else:
        packet = io.BytesIO()
        can = canvas.Canvas(packet, pagesize=A4)
        can.setFont("Helvetica", 10)
        can.drawRightString(19.0 * cm, 27.7 * cm, str(page_num - 1))
        can.showPage()
        can.save()
        
        packet.seek(0)
        num_pdf = pypdf.PdfReader(packet)
        page.merge_page(num_pdf.pages[0])
        writer.add_page(page)

with open(final_pdf_path, "wb") as f_out:
    writer.write(f_out)

print(f"PDF final ABNT gerado com sucesso: {final_pdf_path} ({final_pdf_path.stat().st_size} bytes)")

# Limpeza de arquivos temporários
if raw_pdf_path.exists():
    raw_pdf_path.unlink()
if html_path.exists():
    html_path.unlink()

print("PROCESSO CONCLUÍDO COM SUCESSO!")
print(f"1. DOCX: {docx_path}")
print(f"2. PDF:  {final_pdf_path}")
